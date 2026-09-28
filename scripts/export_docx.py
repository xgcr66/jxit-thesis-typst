"""将 Typst 语义内容导出为可编辑 DOCX。只读取 query 数据，不执行正文代码。"""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

DECLARATION = '本人郑重申明：所呈交的论文（设计）是本人在指导老师的指导下独立进行研究，所取得的研究成果。除了文中特别加以标注引用的内容外，本论文（设计）不包含任何其他个人或集体已经发表或撰写的成果作品。对本文的研究做出重要贡献的个人和集体，均已在文中以明确方式表明。本人完全意识到本申明的法律后果由本人承担。'
AUTHORIZATION = '本学位论文作者完全了解学校有关保留、使用学位论文的规定，同意学校保留并向国家有关部门或机构送交论文的复印件和电子版，允许论文被查阅和借阅。本人授权江西科技学院可以将本论文的全部或部分内容编入有关数据库进行检索，可以采用影印、缩印或扫描等复制手段保存和汇编本学位论文。'


def field(p, instruction, cached=''):
    r = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' ' + instruction + ' '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    text = OxmlElement('w:t'); text.text = str(cached)
    for e in (begin, instr, sep, text, end): r._r.append(e)
    return r


def font(target, cn, size=12, bold=False, latin='Times New Roman'):
    target.font.name = latin
    target.font.size = Pt(size)
    target.font.bold = bold
    target.font.color.rgb = RGBColor(0, 0, 0)
    rpr = target.element.get_or_add_rPr() if hasattr(target, 'element') else target._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), cn)
    rf.set(qn('w:ascii'), latin); rf.set(qn('w:hAnsi'), latin)
    for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'): rf.attrib.pop(qn('w:' + a), None)


def safe_asset(root, name):
    p = (root / name).resolve()
    if not p.is_relative_to(root.resolve()) or not p.is_file():
        raise ValueError(f'图片不存在或超出项目目录：{name}')
    if p.suffix.lower() not in ('.png', '.jpg', '.jpeg'):
        raise ValueError('双格式图片请使用 PNG 或 JPEG；矢量源文件可另行保留。')
    return p


def validate(data, root):
    required = ('meta', 'chapters', 'abstract-cn', 'abstract-en', 'keywords-cn', 'keywords-en', 'references', 'appendix', 'acknowledgements')
    for key in required:
        if key not in data: raise ValueError(f'缺少字段：{key}')
    if not data['chapters']: raise ValueError('至少需要一章正文')
    safe_asset(root, data['meta']['logo'])
    for ch in list(data['chapters']) + [{'blocks': data['appendix']}]:
        has_section = False
        for b in ch['blocks']:
            if b['kind'] not in ('paragraph', 'rich', 'heading', 'figure', 'table', 'code', 'pagebreak'):
                raise ValueError('不支持的内容块：' + b['kind'])
            if b['kind'] in ('paragraph', 'code') and not isinstance(b['text'], str):
                raise ValueError('正文和代码块必须是字符串；任意 Typst 内容不可直接转 Word')
            if b['kind'] == 'heading':
                if b['level'] == 2: has_section = True
                elif b['level'] == 3 and has_section: pass
                else: raise ValueError('三级标题前必须存在二级标题')
            if b['kind'] == 'rich':
                for part in b['parts']:
                    if isinstance(part, str): continue
                    if part.get('kind') != 'citation' or not 1 <= part.get('number', 0) <= len(data['references']):
                        raise ValueError('引用编号应指向现有参考文献')
            if b['kind'] == 'figure':
                safe_asset(root, b['path'])
                if not 0 < b['width'] <= 15: raise ValueError('图宽须为0至15厘米')
            if b['kind'] == 'table':
                n = len(b['header'])
                if not n or any(len(row) != n for row in b['rows']): raise ValueError('表格列数不一致')
                if any(not isinstance(c, str) for row in [b['header'], *b['rows']] for c in row): raise ValueError('单元格必须是字符串')


def build_docx(data, cfg, root, output):
    validate(data, root)
    d = Document()
    d.core_properties.title = data['meta']['title']
    d.core_properties.author = data['meta']['author']
    cn, latin, hei = cfg['body_font'], cfg['latin_font'], cfg['heading_font']
    styles = d.styles
    for name in ('Normal', 'Body Text', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'Header', 'Footer'):
        font(styles[name], cn, latin=latin)
        p = styles[name].paragraph_format
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(20)
        p.widow_control = True
    # 清除默认模板可能携带的标题装饰线。
    for st in styles:
        for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    normal = styles['Normal'].paragraph_format
    normal.first_line_indent = Pt(24); normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, face, size, bold in [(1, hei, 16, True), (2, hei, 12, True), (3, cn, 12, False)]:
        st = styles[f'Heading {level}']; font(st, face, size, bold, latin)
        pf = st.paragraph_format; pf.line_spacing = 1.0; pf.keep_with_next = True
        pf.space_before = Pt(16 if level == 1 else 12); pf.space_after = Pt(16 if level == 1 else 6)
        pf.first_line_indent = Pt(0 if level == 1 else 24)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    font(styles['Title'], hei, 22, True, latin)
    styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles['Title'].paragraph_format.first_line_indent = Pt(0)
    styles['Title'].paragraph_format.line_spacing = 1.0
    for name in ('Caption', 'Header', 'Footer'):
        font(styles[name], cn, 10.5, False, latin)
        pf = styles[name].paragraph_format; pf.first_line_indent = Pt(0); pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing = 1.0
    for level in (1, 2, 3):
        name = f'TOC {level}'
        st = styles[name] if name in styles else styles.add_style(name, 1, builtin=True)
        font(st, cn, 14 if level == 1 else 12, level == 1, latin)
        st.paragraph_format.left_indent = Pt((level-1)*24)
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0 if level == 1 else 3.6)

    def setup_section(s, numbered=True, reset=False, cover=False):
        s.page_width = Cm(21); s.page_height = Cm(29.7)
        s.left_margin = Cm(cfg['left_cm']); s.right_margin = Cm(cfg['right_cm'])
        s.top_margin = Cm(cfg['top_cm'] if cover else max(cfg['top_cm'], cfg['header_cm']+0.7))
        s.bottom_margin = Cm(cfg['bottom_cm'] if cover else max(cfg['bottom_cm'], cfg['footer_cm']+0.7))
        s.gutter = Cm(0); s.header_distance = Cm(cfg['header_cm']); s.footer_distance = Cm(cfg['footer_cm'])
        s.header.is_linked_to_previous = False; s.footer.is_linked_to_previous = False
        hp = s.header.paragraphs[0]; hp.style = 'Header'; hp.text = '' if cover else cfg['header_text']
        fp = s.footer.paragraphs[0]; fp.style = 'Footer'; fp.text = ''
        if numbered: field(fp, 'PAGE', '1')
        if reset:
            el = OxmlElement('w:pgNumType'); el.set(qn('w:start'), '1'); el.set(qn('w:fmt'), 'decimal'); s._sectPr.append(el)

    def section(numbered=True, reset=False):
        s = d.add_section(WD_SECTION_START.NEW_PAGE); setup_section(s, numbered, reset)
        # 分节符独立段落使用最小高度，避免目录末尾挤出空白页。
        marker = d.paragraphs[-1]
        marker.paragraph_format.line_spacing = Pt(1)
        marker.paragraph_format.space_before = marker.paragraph_format.space_after = Pt(0)
        marker.paragraph_format.first_line_indent = Pt(0)
        marker.paragraph_format.keep_with_next = False
        marker.add_run().font.size = Pt(1)
    def para(s='', style=None, indent=True):
        p = d.add_paragraph(s, style)
        if not indent: p.paragraph_format.first_line_indent = Pt(0)
        return p
    def title(text, size=16, face=hei, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = para(indent=False); p.alignment = align
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(16)
        p.paragraph_format.keep_with_next = True; font(p.add_run(text), face, size, True, latin)
        return p
    def space(pt):
        p = para(indent=False); p.paragraph_format.line_spacing = Pt(pt); p.add_run().font.size = Pt(1)
    def caption(text):
        p = para(text, 'Caption', indent=False); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(6)
        return p

    setup_section(d.sections[0], numbered=False, cover=True)
    m = data['meta']
    p=para(f"密级：{m['confidential']}　　学号：{m['student-id']}",indent=False);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    space(26)
    p=para(indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.0
    p.add_run().add_picture(str(safe_asset(root,m['logo'])),width=Cm(3))
    p=para(indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.0
    font(p.add_run(m['school']),hei,28,False,latin)
    space(30)
    p=para('本科生毕业论文（设计）','Title',False)
    space(28)
    title(m['title'],18)
    space(28)
    for label,key in [('学　　院','college'),('专　　业','major'),('班　　级','class'),('学　　号','student-id'),('姓　　名','author'),('指导教师1','supervisor'),('指导教师2','second-supervisor')]:
        p=para(f'{label}：　{m[key]}',indent=False);p.paragraph_format.left_indent=Cm(2.2);p.paragraph_format.space_after=Pt(7)
    space(15)
    p=para(m['date'],indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER

    section(numbered=False)
    title('学士学位论文原创性声明');para(DECLARATION)
    p=para('学位论文作者签名（手写）：____________\n签字日期：　　年　　月　　日',indent=False);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.space_before=Pt(12)
    title('学位论文版权使用授权书');para(AUTHORIZATION)
    for s in ['本学位论文属于：','保密□，在　　　年解密后适用本授权书。','不保密□。','（请在以上相应方框内打“√”）']:para(s)
    p=para('作者签名（手写）：________　指导教师签名（手写）：________\n签字日期：　　年　月　日　　签字日期：　　年　月　日',indent=False)
    p.paragraph_format.space_before=Pt(14)
    for r in p.runs: font(r,cn,10.5,False,latin)

    section(reset=True)
    for lang,label in [('cn','摘　要'),('en','Abstract')]:
        if lang=='en':d.add_page_break()
        title(label,face=hei if lang=='cn' else latin)
        for s in data['abstract-'+lang]:para(s)
        p=para(indent=False);p.paragraph_format.space_before=Pt(20)
        p.add_run('关键词：' if lang=='cn' else 'Key words: ').bold=True
        p.add_run(('；' if lang=='cn' else '; ').join(data['keywords-'+lang]))
    section(reset=True);title('目　录')
    toc=para(indent=False);field(toc, 'TOC \\o "1-3" \\h \\z \\u', '请在 Word 中右键更新域，选择更新整个目录。')
    section(reset=True)

    def blocks(items, ch):
        sec=sub=fig=tab=0
        for b in items:
            kind=b['kind']
            if kind=='paragraph':para(b['text'])
            elif kind=='rich':
                p=para()
                for part in b['parts']:
                    r=p.add_run(part if isinstance(part,str) else '['+str(part['number'])+']')
                    if not isinstance(part,str):r.font.superscript=True
            elif kind=='heading':
                if b['level']==2:sec+=1;sub=0;num=f'{ch}.{sec}'
                else:sub+=1;num=f'{ch}.{sec}.{sub}'
                para(num+'　'+b['text'],f"Heading {b['level']}")
            elif kind=='figure':
                fig+=1;p=para(indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1.0
                p.add_run().add_picture(str(safe_asset(root,b['path'])),width=Cm(b['width']))
                caption(f'图{ch}.{fig} '+b['caption'])
            elif kind=='table':
                tab+=1;p=caption(f'表{ch}.{tab} '+b['caption']);p.paragraph_format.keep_with_next=True
                t=d.add_table(rows=1,cols=len(b['header']));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
                n=len(b['header']);weights=([.08,.18,.26,.29,.19] if n==5 else [.23,.18,.22,.37] if n==4 else [1/n]*n)
                for col,w in zip(t.columns,weights):col.width=Cm(15*w)
                for rownum,vals in enumerate([b['header'],*b['rows']]):
                    cells=t.rows[0].cells if rownum==0 else t.add_row().cells
                    trpr=cells[0]._tc.getparent().get_or_add_trPr()
                    trpr.append(OxmlElement('w:cantSplit'))
                    if rownum==0:trpr.append(OxmlElement('w:tblHeader'))
                    for j,(cell,val) in enumerate(zip(cells,vals)):
                        cell.width=Cm(15*weights[j]);p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(16)
                        p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(4)
                        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if rownum==0 else WD_ALIGN_PARAGRAPH.LEFT
                        font(p.add_run(val),cn,10.5,rownum==0,latin)
                        borders=OxmlElement('w:tcBorders')
                        for edge in ('top','bottom','left','right'):
                            e=OxmlElement('w:'+edge)
                            visible=(edge=='top' and rownum==0) or (edge=='bottom' and rownum in (0,len(b['rows'])))
                            e.set(qn('w:val'),'single' if visible else 'nil');e.set(qn('w:sz'),'6' if rownum==0 and edge=='top' or rownum==len(b['rows']) else '4');e.set(qn('w:color'),'000000');borders.append(e)
                        cell._tc.get_or_add_tcPr().append(borders)
                para(indent=False)
            elif kind=='code':
                p=para(indent=False);p.paragraph_format.line_spacing=Pt(12);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                font(p.add_run(b['text']),cn,9,False,'Consolas')
            elif kind=='pagebreak':d.add_page_break()
    for i,ch in enumerate(data['chapters'],1):
        if i>1:d.add_page_break()
        para(f"第{i}章　{ch['title']}",'Heading 1',False)
        blocks(ch['blocks'],i)
    d.add_page_break()
    p=para('参考文献','Heading 1',False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:font(r,hei,14,True,latin)
    for i,s in enumerate(data['references'],1):
        p=para(f'[{i}] {s}',indent=False);p.paragraph_format.line_spacing=Pt(16)
    if data['appendix']:
        d.add_page_break();para('附　录','Heading 1',False);blocks(data['appendix'],'A')
    if data['acknowledgements']:
        d.add_page_break();para('致　谢','Heading 1',False)
        for s in data['acknowledgements']:
            p=para(s)
            for r in p.runs:font(r,cfg['ack_font'],12,False,latin)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');d.settings.element.append(update)
    d.save(output)
    return output
