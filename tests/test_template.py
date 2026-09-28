import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from export_docx import validate,build_docx

class TemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data=json.loads((ROOT/'build/content.json').read_text(encoding='utf-8'))
        cls.cfg=json.loads((ROOT/'layout.json').read_text(encoding='utf-8-sig'))
        cls.temp=tempfile.TemporaryDirectory()
        cls.path=Path(cls.temp.name)/'test.docx'
        build_docx(cls.data,cls.cfg,ROOT,cls.path)
        cls.doc=Document(cls.path)
    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()
    def test_content_valid(self):validate(self.data,ROOT)
    def test_six_chapters(self):self.assertEqual(len(self.data['chapters']),6)
    def test_a4_and_margins(self):
        s=self.doc.sections[-1]
        self.assertAlmostEqual(s.page_width.cm,21,places=2)
        self.assertAlmostEqual(s.page_height.cm,29.7,places=2)
        self.assertAlmostEqual(s.left_margin.cm,3,places=2)
        self.assertAlmostEqual(s.header_distance.cm,2.8,places=2)
        self.assertAlmostEqual(s.footer_distance.cm,2.8,places=2)
        self.assertAlmostEqual(s.top_margin.cm,3.5,places=2)
    def test_body_style(self):
        s=self.doc.styles['Normal'];self.assertEqual(s.font.size.pt,12)
        self.assertEqual(s.paragraph_format.line_spacing.pt,20)
        self.assertEqual(s.paragraph_format.first_line_indent.pt,24)
    def test_editable_tables_and_all_cells(self):
        expected=[b for c in self.data['chapters'] for b in c['blocks'] if b['kind']=='table']
        self.assertEqual(len(self.doc.tables),len(expected))
        for table,source in zip(self.doc.tables,expected):
            self.assertEqual([[c.text for c in r.cells] for r in table.rows],[source['header'],*source['rows']])
            self.assertTrue(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader'))
            self.assertTrue(all(r._tr.xpath('./w:trPr/w:cantSplit') for r in table.rows))
    def test_images_have_automatic_line_height(self):
        count=0
        for p in self.doc.paragraphs:
            if p._p.xpath('.//w:drawing'):
                count+=1
                self.assertEqual(p.paragraph_format.line_spacing,1.0)
        self.assertEqual(count,2)
    def test_word_fields_and_section_resets(self):
        xml=self.doc.element.xml
        self.assertIn('TOC \\o "1-3"',xml)
        starts=[s._sectPr.find(qn('w:pgNumType')) for s in self.doc.sections]
        self.assertEqual(sum(x is not None and x.get(qn('w:start'))=='1' for x in starts),3)
    def test_titles_no_decoration(self):
        self.assertFalse(self.doc.styles['Title'].element.xpath('.//w:pBdr'))
        self.assertFalse(self.doc.styles['Title'].font.underline)
    def test_all_paragraph_content_preserved(self):
        text='\n'.join(p.text for p in self.doc.paragraphs)
        for ch in self.data['chapters']:
            self.assertIn(ch['title'],text)
            for b in ch['blocks']:
                if b['kind'] in ('paragraph','code'):self.assertIn(b['text'],text)
    def test_unknown_block_rejected(self):
        d=copy.deepcopy(self.data);d['chapters'][0]['blocks'].append({'kind':'custom-typst'})
        with self.assertRaises(ValueError):validate(d,ROOT)
    def test_missing_asset_rejected(self):
        d=copy.deepcopy(self.data);d['meta']['logo']='assets/no-such-image.png'
        with self.assertRaises(ValueError):validate(d,ROOT)
    def test_external_asset_rejected(self):
        d=copy.deepcopy(self.data);d['meta']['logo']='../external.jpg'
        with self.assertRaises(ValueError):validate(d,ROOT)
    def test_invalid_citation_rejected(self):
        d=copy.deepcopy(self.data);d['chapters'][0]['blocks'].append({'kind':'rich','parts':[{'kind':'citation','number':99}]})
        with self.assertRaises(ValueError):validate(d,ROOT)
    def test_bad_table_rejected(self):
        d=copy.deepcopy(self.data);d['chapters'][0]['blocks'].append({'kind':'table','header':['a','b'],'rows':[['a']]})
        with self.assertRaises(ValueError):validate(d,ROOT)
    def test_bad_heading_order_rejected(self):
        d=copy.deepcopy(self.data);d['chapters'][0]['blocks'].insert(0,{'kind':'heading','level':3,'text':'bad'})
        with self.assertRaises(ValueError):validate(d,ROOT)
    def test_optional_sections_can_be_hidden(self):
        d=copy.deepcopy(self.data);d['appendix']=[];d['acknowledgements']=[]
        p=Path(self.temp.name)/'optional.docx';build_docx(d,self.cfg,ROOT,p)
        texts=[p.text for p in Document(p).paragraphs]
        self.assertNotIn('附　录',texts);self.assertNotIn('致　谢',texts)

if __name__=='__main__':unittest.main()
